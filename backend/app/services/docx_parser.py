import io
from docx import Document

def process_docx(file_bytes: bytes) -> dict:
    doc = Document(io.BytesIO(file_bytes))
    structured_content = []

    # 1. Extract Paragraphs with Structure (Headings vs Normal text)
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        
        style = para.style.name.lower()
        if 'heading' in style:
            structured_content.append({"type": "heading", "content": text, "level": style})
        else:
            structured_content.append({"type": "paragraph", "content": text})

    # 2. Extract Tables
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        if table_data:
            structured_content.append({"type": "table", "content": table_data})

    return {"data": structured_content, "low_quality": False}